import re
from pathlib import Path

import pandas as pd

from config import DATA_DIR
from src.schema import get_standard_name


def normalize_unit_name(unit_name: str) -> str:
    if unit_name is None:
        raise ValueError("A unidade não foi informada.")

    unit_name = str(unit_name).strip().upper()

    if not unit_name:
        raise ValueError("A unidade está vazia.")

    return unit_name


def get_unit_file_candidates(unit_name: str) -> list[Path]:
    """
    Prioriza CSV e depois XLSX.
    """
    unit_name = normalize_unit_name(unit_name)

    return [
        DATA_DIR / f"{unit_name}_data.csv",
        DATA_DIR / f"{unit_name}_data.xlsx",
    ]


def list_available_unit_files() -> list[str]:
    """
    Lista arquivos disponíveis na pasta data nos padrões:
    - *_data.csv
    - *_data.xlsx
    """
    if not DATA_DIR.exists():
        return []

    files = []
    files.extend([file.name for file in DATA_DIR.glob("*_data.csv")])
    files.extend([file.name for file in DATA_DIR.glob("*_data.xlsx")])

    return sorted(files)


def read_csv_safely(file_path: Path) -> pd.DataFrame:
    """
    Tenta ler CSV com combinações comuns de separador e encoding.

    pd.read_csv não levanta erro quando o separador está errado: devolve uma
    única coluna com a linha inteira. Por isso cada tentativa é validada pelo
    número de colunas obtidas, e não apenas pela ausência de exceção.

    Tudo é lido como texto: sem isso o código de cliente 00123 vira o número
    123. Datas e volume já são convertidos a partir de texto adiante.
    """
    attempts = [
        {"sep": ";", "encoding": "utf-8-sig"},
        {"sep": ",", "encoding": "utf-8-sig"},
        {"sep": ";", "encoding": "latin1"},
        {"sep": ",", "encoding": "latin1"},
        {"sep": ";", "encoding": "cp1252"},
        {"sep": ",", "encoding": "cp1252"},
    ]

    last_error = None
    tentativas_uma_coluna = []

    for params in attempts:
        try:
            df = pd.read_csv(file_path, low_memory=False, dtype=str, **params)
        except Exception as exc:
            last_error = exc
            continue

        if df.shape[1] > 1:
            return df

        primeira_coluna = str(df.columns[0]) if df.shape[1] else ""
        tentativas_uma_coluna.append(
            f"sep='{params['sep']}' encoding='{params['encoding']}' "
            f"-> 1 coluna ('{primeira_coluna[:60]}')"
        )

    if tentativas_uma_coluna:
        raise ValueError(
            f"Erro ao ler o CSV '{file_path.name}'. "
            f"Nenhum separador testado produziu mais de uma coluna. "
            f"Tentativas: {' | '.join(tentativas_uma_coluna)}"
        )

    raise ValueError(
        f"Erro ao ler o CSV '{file_path.name}'. "
        f"Não foi possível interpretar o arquivo com os formatos testados. "
        f"Último erro: {last_error}"
    )


def _codigo_como_texto(celula) -> str | None:
    """
    Texto do codigo do cliente como o operador o ve na celula.

    Numero com formato so de zeros (00000) e um codigo com zeros a esquerda
    guardado como numero: o formato devolve os zeros. Qualquer outro formato
    fica com os digitos do numero, sem zeros a esquerda.
    """
    valor = celula.value
    if valor is None:
        return None
    if isinstance(valor, bool) or not isinstance(valor, (int, float)):
        return str(valor)
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = str(valor)
    formato = str(celula.number_format or "")
    if isinstance(valor, int) and valor >= 0 and re.fullmatch(r"0+", formato):
        texto = texto.zfill(len(formato))
    return texto


def read_excel_safely(file_path: Path) -> pd.DataFrame:
    """
    Le a primeira aba com Cod_Cliente como texto.

    pd.read_excel converte o texto "00123" no numero 123, e o cliente se
    funde com o 123. Aqui a coluna do codigo (reconhecida pelos mesmos aliases
    da padronizacao) sai como texto; as demais ficam com o tipo da celula.
    """
    from openpyxl import load_workbook

    wb = load_workbook(file_path, read_only=True, data_only=True)
    try:
        linhas = wb.worksheets[0].iter_rows()
        cabecalho = [
            str(c.value) if c.value is not None else f"Unnamed: {i}"
            for i, c in enumerate(next(linhas, ()))
        ]
        largura = len(cabecalho)
        colunas_codigo = {
            i for i, nome in enumerate(cabecalho) if get_standard_name(nome) == "Cod_Cliente"
        }

        dados = []
        for linha in linhas:
            valores = [None] * largura
            for i, celula in enumerate(linha[:largura]):
                valores[i] = (
                    _codigo_como_texto(celula) if i in colunas_codigo else celula.value
                )
            dados.append(valores)
    finally:
        wb.close()

    # Linhas vazias no fim da aba sao formatacao, nao registros.
    while dados and all(v is None for v in dados[-1]):
        dados.pop()

    return pd.DataFrame(dados, columns=cabecalho)


def carregar_arquivo(caminho: Path) -> pd.DataFrame:
    """
    Carrega uma base a partir de um caminho informado pelo operador.

    Aceita qualquer lugar do disco ou da rede (V:, G:, N:, P:), sem a convenção
    de nome {UNIDADE}_data.csv exigida por load_unit_file. CSV passa por
    read_csv_safely, que é quem resolve separador e encoding das bases vindas
    do file server.
    """
    caminho = Path(caminho)

    if not caminho.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {caminho}")

    if not caminho.is_file():
        raise ValueError(f"O caminho informado não é um arquivo: {caminho}")

    sufixo = caminho.suffix.lower()

    if sufixo == ".csv":
        return read_csv_safely(caminho)

    if sufixo in {".xlsx", ".xlsm"}:
        return read_excel_safely(caminho)

    raise ValueError(
        f"Formato não suportado: '{sufixo}'. Esperado .csv, .xlsx ou .xlsm."
    )


def load_unit_file(unit_name: str) -> pd.DataFrame:
    """
    Carrega o arquivo da unidade a partir da pasta data.
    Prioridade:
    1. {UNIDADE}_data.csv
    2. {UNIDADE}_data.xlsx
    """
    candidates = get_unit_file_candidates(unit_name)

    selected_file = None
    for file_path in candidates:
        if file_path.exists():
            selected_file = file_path
            break

    if selected_file is None:
        available_files = list_available_unit_files()
        available_text = ", ".join(available_files) if available_files else "nenhum arquivo encontrado"

        expected_names = ", ".join([path.name for path in candidates])

        raise FileNotFoundError(
            f"Arquivo não encontrado para a unidade '{normalize_unit_name(unit_name)}'. "
            f"Nomes esperados: {expected_names}. "
            f"Pasta pesquisada: {DATA_DIR}. "
            f"Arquivos disponíveis: {available_text}"
        )

    try:
        if selected_file.suffix.lower() == ".csv":
            return read_csv_safely(selected_file)

        if selected_file.suffix.lower() == ".xlsx":
            return read_excel_safely(selected_file)

        raise ValueError(f"Formato não suportado: {selected_file.suffix}")

    except Exception as exc:
        raise ValueError(
            f"Erro ao ler o arquivo '{selected_file.name}': {exc}"
        ) from exc