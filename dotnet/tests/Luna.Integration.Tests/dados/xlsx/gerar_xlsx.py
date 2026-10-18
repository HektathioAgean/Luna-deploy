"""
Planilhas sintéticas para a paridade da leitura de Excel (PR6b) e o que a
referência Python lê de cada uma.

    .venv\\Scripts\\python dotnet\\tests\\Luna.Integration.Tests\\dados\\xlsx\\gerar_xlsx.py

O XML é escrito à mão (o openpyxl não grava strings inline nem valores em
cache de fórmula) e o zip é determinístico. Para cada planilha, grava em
esperado.json:
- "linhas": cada célula de wb.worksheets[0].iter_rows() do openpyxl
  (read_only=True, data_only=True): tipo, valor e number_format;
- "tabela": o DataFrame de read_excel_safely (src/data_loader.py).
Para as bases de ponta a ponta, grava também as saídas de executar().

Nenhum código do motor muda: este arquivo só usa a referência.
"""

from __future__ import annotations

import datetime as dt
import json
from html import escape
import sys
import warnings
import zipfile
from pathlib import Path

PASTA = Path(__file__).resolve().parent
RAIZ = PASTA.parents[4]
sys.path.insert(0, str(RAIZ))

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from src.data_loader import read_excel_safely  # noqa: E402
from src.pipeline import Parametros, executar  # noqa: E402

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_MAIN = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
CT_MACRO = "application/vnd.ms-excel.sheet.macroEnabled.main+xml"
CT_SHEET = "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"
CT_CHART = "application/vnd.openxmlformats-officedocument.spreadsheetml.chartsheet+xml"
CT_STYLES = "application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"
CT_SST = "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
DATA_ZIP = (2026, 1, 1, 0, 0, 0)


# -- Montagem do pacote ---------------------------------------------------------

def gravar_xlsx(nome, abas, estilos=None, strings=None, data1904=False, macro=False):
    """abas: lista de (nome_da_aba, tipo, xml, oculta); tipo 'sheet' ou 'chart'."""
    partes = {}
    overrides = [("/xl/workbook.xml", CT_MACRO if macro else CT_MAIN)]
    rels = []
    sheets = []
    for i, (titulo, tipo, xml, oculta) in enumerate(abas, start=1):
        pasta = "worksheets" if tipo == "sheet" else "chartsheets"
        alvo = f"{pasta}/sheet{i}.xml"
        partes[f"xl/{alvo}"] = xml
        if tipo == "chart":
            partes[f"xl/{pasta}/_rels/sheet{i}.xml.rels"] = (
                f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="{NS_PKG}"/>')
        overrides.append((f"/xl/{alvo}", CT_SHEET if tipo == "sheet" else CT_CHART))
        rels.append((f"rId{i}", f"{REL}/{'worksheet' if tipo == 'sheet' else 'chartsheet'}", alvo))
        estado = ' state="hidden"' if oculta else ""
        sheets.append(f'<sheet name="{titulo}" sheetId="{i}"{estado} r:id="rId{i}"/>')
    n = len(abas)
    if estilos is not None:
        partes["xl/styles.xml"] = estilos
        overrides.append(("/xl/styles.xml", CT_STYLES))
        rels.append((f"rId{n + 1}", f"{REL}/styles", "styles.xml"))
    if strings is not None:
        itens = "".join(f"<si>{s}</si>" for s in strings)
        partes["xl/sharedStrings.xml"] = (
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<sst xmlns="{NS}" count="{len(strings)}" uniqueCount="{len(strings)}">{itens}</sst>')
        overrides.append(("/xl/sharedStrings.xml", CT_SST))
        rels.append((f"rId{n + 2}", f"{REL}/sharedStrings", "sharedStrings.xml"))
    pr = '<workbookPr date1904="1"/>' if data1904 else "<workbookPr/>"
    partes["xl/workbook.xml"] = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<workbook xmlns="{NS}" xmlns:r="{NS_R}">{pr}<sheets>{"".join(sheets)}</sheets></workbook>')
    partes["xl/_rels/workbook.xml.rels"] = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="{NS_PKG}">'
        + "".join(f'<Relationship Id="{i}" Type="{t}" Target="{a}"/>' for i, t, a in rels)
        + "</Relationships>")
    partes["_rels/.rels"] = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="{NS_PKG}">'
        f'<Relationship Id="rId1" Type="{REL}/officeDocument" Target="xl/workbook.xml"/></Relationships>')
    partes["[Content_Types].xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        + "".join(f'<Override PartName="{p}" ContentType="{c}"/>' for p, c in overrides)
        + "</Types>")
    caminho = PASTA / nome
    with zipfile.ZipFile(caminho, "w", zipfile.ZIP_DEFLATED) as z:
        for parte in sorted(partes, key=lambda p: (p != "[Content_Types].xml", p)):
            info = zipfile.ZipInfo(parte, DATA_ZIP)
            info.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(info, partes[parte].encode("utf-8"))
    return caminho


def planilha(linhas, dimensao=None):
    dim = f'<dimension ref="{dimensao}"/>' if dimensao else ""
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<worksheet xmlns="{NS}" xmlns:r="{NS_R}">{dim}<sheetData>{"".join(linhas)}</sheetData></worksheet>')


def estilos(formatos, xfs):
    """formatos: {id: código}; xfs: lista de numFmtId (um por estilo de célula)."""
    fmts = "".join(f'<numFmt numFmtId="{i}" formatCode="{escape(c, quote=True)}"/>' for i, c in formatos.items())
    cells = "".join(f'<xf numFmtId="{n}" fontId="0" fillId="0" borderId="0" xfId="0"/>' for n in xfs)
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><styleSheet xmlns="{NS}">'
            f'<numFmts count="{len(formatos)}">{fmts}</numFmts>'
            '<fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>'
            '<fills count="1"><fill><patternFill patternType="none"/></fill></fills>'
            '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
            '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
            f'<cellXfs count="{len(xfs)}">{cells}</cellXfs></styleSheet>')


def c(ref, valor=None, t=None, s=None, f=None, inline=None):
    attrs = f' r="{ref}"' if ref else ""
    if t:
        attrs += f' t="{t}"'
    if s is not None:
        attrs += f' s="{s}"'
    corpo = ""
    if f is not None:
        corpo += f"<f>{f}</f>"
    if inline is not None:
        corpo += f"<is>{inline}</is>"
    if valor is not None:
        corpo += f"<v>{valor}</v>"
    return f"<c{attrs}>{corpo}</c>" if corpo else f"<c{attrs}/>"


def linha(r, *celulas, extra=""):
    attr = f' r="{r}"' if r is not None else ""
    return f"<row{attr}{extra}>{''.join(celulas)}</row>"


def serial(data):
    return repr((data - dt.datetime(1899, 12, 30)).total_seconds() / 86400)


# -- Planilhas -------------------------------------------------------------------

# Estilos de célula usados nas planilhas "tipos": índice -> numFmtId.
FORMATOS = {
    164: "00000", 165: "dd/mm/yyyy hh:mm", 166: "0.00", 167: "[h]:mm:ss",
    169: "0000;-0000", 170: '"Cód "0', 171: "[Red]0",
}
XFS = [0, 14, 22, 164, 165, 20, 46, 166, 1, 27, 169, 170, 45, 49, 10, 167, 171]
#      0  1   2   3    4    5   6   7    8  9   10   11   12  13  14  15   16


def planilha_tipos():
    sst = [
        "<t>Código do cliente</t>", "<t>cliente</t>", "<t>00123</t>",
        '<r><t>C</t></r><r><rPr><b/></rPr><t xml:space="preserve">0 1</t></r>',
        "<t>a_x005F_x0041_b</t>", "<t>Kanji</t><rPh sb=\"0\" eb=\"1\"><t>kana</t></rPh>",
        '<t xml:space="preserve">  E1 </t>', "<t>texto</t>",
    ]
    linhas = [
        linha(1, c("A1", 0, "s"), c("B1", t="inlineStr", inline="<t>Valor</t>"), c("C1", 2026),
              c("E1", 1, "s"), c("F1", 1, "b"), c("G1", 1.5)),
        linha(2, c("A2", 2, "s"), c("B2", 1.5), c("E2", 6, "s"), c("F2", 7, "s")),
        linha(3, c("A3", 123, s=3), c("B3", 59, s=1), c("E3", 99, s=3)),
        linha(4, c("A4", 123), c("B4", 60, s=1)),
        linha(5, c("A5", 12.5), c("B5", 61, s=1)),
        linha(6, c("A6", "123.0", s=3), c("B6", 0, s=1)),
        linha(7, c("A7", -5, s=3), c("B7", 0.375, s=5)),
        linha(8, c("A8", 1, "b"), c("B8", 1.5, s=6)),
        linha(9, c("A9", 46083, s=1), c("B9", 46083.3819444, s=2)),
        linha(10, c("A10", "#N/A", "e"), c("B10", 46083.123456789, s=4)),
        linha(11, c("A11", "ABC", "str", f="CONCAT(&quot;A&quot;,&quot;BC&quot;)"), c("B11", -1, s=1)),
        linha(12, c("A12", t="inlineStr", inline='<r><t>C</t></r><r><t>01</t></r>'), c("B12", 3000000, s=1)),
        linha(13, c("A13", 7, s=8), c("B13", "2026-03-02T09:15:30.5", "d")),
        linha(14, c("A14", 7, s=10), c("B14", "2026-03-02", "d")),
        linha(15, c("A15", "1E3", s=3), c("B15", "09:15", "d")),
        linha(16, c("A16", 0.1), c("B16", "PT2H30M", "d")),
        linha(17, c("A17", "1E-05"), c("B17", 0.5, s=12)),
        linha(18, c("A18", 12345678901234567), c("B18", 0.999999999, s=1)),
        linha(19, c("A19", s=3), c("B19", 1.5, s=15)),
        linha(20, c("A20", 5, s=9), c("B20", 1.23456789, s=6)),
        linha(21, c("A21", 42, s=3, f="40+2"), c("B21", 0, "b")),
        linha(22, c("A22", f="A1"), c("B22", "#DIV/0!", "e")),
        linha(23, c("A23", 4, "s"), c("B23", 46083, s=13)),
        linha(24, c("A24", 5, "s"), c("B24", 0.25, s=14)),
        linha(25, c("A25", 7, s=11), c("B25", 2.0000000001, s=5)),
        linha(26, c("A26", 7, s=16), c("B26", 1, s=7)),
        linha(27, c("A27", 3, s=1), c("B27", "", "str", f="&quot;&quot;")),
        linha(28, c("A28", "0.5", s=3), c("B28", 1, s=0)),
        # 549348.5 ms exatos: o round half-even do Python dá 549348.
        linha(29, c("A29", 29), c("B29", "1.0063582002314815", s=2)),
    ]
    return gravar_xlsx("tipos.xlsx", [("Dados", "sheet", planilha(linhas, "A1:G29"), False)],
                       estilos(FORMATOS, XFS), sst)


def planilha_esparsa():
    linhas = [
        linha(1, c("A1", "<t>Cod_Cliente</t>", "inlineStr"), c("C1", t="inlineStr", inline="<t>Chegou_em</t>"),
              c("E1", t="inlineStr", inline="<t>Finalizada_em</t>")),
        linha(2, c(None, 1), c(None, t="inlineStr", inline="<t>x</t>"), c(None, 3)),
        linha(4, c("A4", 4)),
        linha("5.0", c("B5", 5)),
        linha(6),
        linha(7, c("A7", 7), c("F7", 99)),
        linha(None, c(None, 8)),
        linha(9, c("A9", s=0), c("C9", s=0)),
    ]
    return gravar_xlsx("esparsa.xlsx", [("S", "sheet", planilha(linhas, "A1:E9"), False)], estilos({}, [0]))


def planilha_sem_dimensao():
    linhas = [
        linha(1, c("A1", t="inlineStr", inline="<t>Cod_Cliente</t>"), c("B1", t="inlineStr", inline="<t>X</t>")),
        linha(2, c("A2", 1), c("D2", 4)),
        linha(4, c("B4", 2)),
        linha(5, c("A5", s=0)),
    ]
    return gravar_xlsx("sem_dimensao.xlsx", [("S", "sheet", planilha(linhas), False)], estilos({}, [0]))


def planilha_dimensao_errada():
    linhas = [linha(r, *(c(f"{col}{r}", r * 10 + i) for i, col in enumerate("ABCD"))) for r in range(1, 6)]
    return gravar_xlsx("dimensao_errada.xlsx", [("S", "sheet", planilha(linhas, "A1:B2"), False)], estilos({}, [0]))


def planilha_dimensao_maior():
    """<dimension> com mais linhas que o arquivo: o openpyxl não completa as do fim."""
    linhas = [linha(1, c("A1", t="inlineStr", inline="<t>Cod_Cliente</t>"), c("B1", 1)), linha(2, c("A2", 2))]
    return gravar_xlsx("dimensao_maior.xlsx", [("S", "sheet", planilha(linhas, "A1:B6"), False)], estilos({}, [0]))


def planilha_1904():
    linhas = [linha(1, c("A1", t="inlineStr", inline="<t>Data</t>"))] + [
        linha(i + 2, c(f"A{i + 2}", v, s=1)) for i, v in enumerate([0, 1, 59, 60, 61, 0.5, 44621.25, -1])]
    return gravar_xlsx("epoca1904.xlsx", [("S", "sheet", planilha(linhas, f"A1:A{len(linhas)}"), False)],
                       estilos({}, [0, 22]), data1904=True)


def planilha_abas():
    primeira = planilha([linha(1, c("A1", t="inlineStr", inline="<t>Cod_Cliente</t>")), linha(2, c("A2", 1))], "A1:A2")
    segunda = planilha([linha(1, c("A1", t="inlineStr", inline="<t>Outra</t>"))], "A1")
    grafico = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><chartsheet xmlns="{NS}" xmlns:r="{NS_R}">'
               '<sheetViews><sheetView workbookViewId="0"/></sheetViews></chartsheet>')
    # Ordem do workbook: gráfico, aba oculta com dados, aba visível.
    return gravar_xlsx("abas.xlsx", [("Grafico", "chart", grafico, False), ("Oculta", "sheet", primeira, True),
                                     ("Visivel", "sheet", segunda, False)], estilos({}, [0]))


def planilha_sem_estilos():
    linhas = [linha(1, c("A1", t="inlineStr", inline="<t>Cod_Cliente</t>"), c("B1", t="inlineStr", inline="<t>D</t>")),
              linha(2, c("A2", 123), c("B2", 46083.5))]
    return gravar_xlsx("sem_estilos.xlsx", [("S", "sheet", planilha(linhas, "A1:B2"), False)])


def planilha_macro():
    linhas = [linha(1, c("A1", t="inlineStr", inline="<t>Cod_Cliente</t>"), c("B1", t="inlineStr", inline="<t>N</t>")),
              linha(2, c("A2", 7, s=1), c("B2", 2))]
    return gravar_xlsx("macro.xlsm", [("S", "sheet", planilha(linhas, "A1:B2"), False)],
                       estilos({164: "00000"}, [0, 164]), macro=True)


def planilha_base_sintetica():
    """tests/dados/TESTE_data.csv em Excel: códigos e textos como strings, datas como seriais com formato de data."""
    csv = pd.read_csv(RAIZ / "tests" / "dados" / "TESTE_data.csv", sep=";", dtype=str, keep_default_na=False)
    sst, indice = [], {}

    def s(texto):
        if texto not in indice:
            indice[texto] = len(sst)
            sst.append(f'<t xml:space="preserve">{texto}</t>')
        return indice[texto]

    colunas = list(csv.columns)
    letras = "ABCDEFGHIJ"
    linhas = [linha(1, *(c(f"{letras[j]}1", s(n), "s") for j, n in enumerate(colunas)))]
    for i, registro in enumerate(csv.itertuples(index=False), start=2):
        celulas = []
        for j, (coluna, valor) in enumerate(zip(colunas, registro)):
            ref = f"{letras[j]}{i}"
            if valor == "":
                continue
            if coluna in ("Chegou_em", "Finalizada_em"):
                try:
                    celulas.append(c(ref, serial(dt.datetime.fromisoformat(valor)), s=2))
                    continue
                except ValueError:
                    pass
            celulas.append(c(ref, s(valor), "s"))
        linhas.append(linha(i, *celulas))
    return gravar_xlsx("base_sintetica.xlsx", [("Base", "sheet", planilha(linhas, f"A1:F{len(linhas)}"), False)],
                       estilos({}, [0, 14, 22]), sst)


def planilha_codigos():
    """Códigos numéricos com e sem formato 00000, texto com zeros e datas nativas e em texto BR."""
    inicio = dt.datetime(2026, 3, 2, 8, 0)
    linhas = [linha(1, c("A1", t="inlineStr", inline="<t>Código do Cliente</t>"),
                    c("B1", t="inlineStr", inline="<t>Chegou em</t>"),
                    c("C1", t="inlineStr", inline="<t>Finalizada_em</t>"))]
    codigos = [(123, 3), (123, 0), ("00123", None), (45, 3), ("45", None)]
    r = 2
    for k, (codigo, estilo) in enumerate(codigos):
        for dia in range(5):
            chegou = inicio + dt.timedelta(days=dia, minutes=10 * k)
            fim = chegou + dt.timedelta(minutes=10 + dia + k)
            a = (c(f"A{r}", codigo, s=estilo) if isinstance(codigo, int)
                 else c(f"A{r}", t="inlineStr", inline=f"<t>{codigo}</t>"))
            if dia % 2:
                b = c(f"B{r}", t="inlineStr", inline=f"<t>{chegou:%d/%m/%Y %H:%M}</t>")
                cc = c(f"C{r}", t="inlineStr", inline=f"<t>{fim:%d/%m/%Y %H:%M}</t>")
            else:
                b, cc = c(f"B{r}", serial(chegou), s=2), c(f"C{r}", serial(fim), s=2)
            linhas.append(linha(r, a, b, cc))
            r += 1
    return gravar_xlsx("codigos.xlsx", [("Base", "sheet", planilha(linhas, f"A1:C{r - 1}"), False)],
                       estilos({164: "00000"}, [0, 14, 22, 164]))


# -- Referência ------------------------------------------------------------------

def canon(v):
    if v is None or v is pd.NaT:
        return ["nulo"]
    if isinstance(v, (bool, np.bool_)):
        return ["bool", str(bool(v))]
    if isinstance(v, (int, np.integer)):
        return ["int", str(int(v))]
    if isinstance(v, (float, np.floating)):
        return ["nulo"] if np.isnan(v) else ["float", repr(float(v))]
    if isinstance(v, str):
        return ["str", v]
    if isinstance(v, pd.Timestamp):
        return ["nulo"] if pd.isna(v) else ["datetime", v.strftime("%Y-%m-%dT%H:%M:%S.%f")]
    if isinstance(v, dt.datetime):
        return ["datetime", v.strftime("%Y-%m-%dT%H:%M:%S.%f")]
    if isinstance(v, dt.date):
        return ["date", v.isoformat()]
    if isinstance(v, dt.time):
        return ["time", v.strftime("%H:%M:%S.%f")]
    if isinstance(v, pd.Timedelta):
        return ["nulo"] if pd.isna(v) else ["timedelta", str(v.value // 1000)]
    if isinstance(v, dt.timedelta):
        return ["timedelta", str((v.days * 86400 + v.seconds) * 1_000_000 + v.microseconds)]
    raise TypeError(type(v))


def referencia(caminho):
    wb = load_workbook(caminho, read_only=True, data_only=True)
    linhas = [[canon(cel.value) + [cel.number_format or ""] for cel in r]
              for r in wb.worksheets[0].iter_rows()]
    wb.close()
    df = read_excel_safely(caminho)
    tabela = {"colunas": [str(col) for col in df.columns],
              "linhas": [[canon(v) for v in r] for r in df.astype(object).values.tolist()]}
    return {"linhas": linhas, "tabela": tabela}


def saidas_do_processamento(caminho, nome):
    r = executar(Parametros(caminho=str(caminho)), usar_cache=False)
    r["medianas"].to_csv(PASTA / f"{nome}_medianas.csv", sep=";", index=False, lineterminator="\n")
    r["janelas_entrega"].to_csv(PASTA / f"{nome}_janelas.csv", sep=";", index=False, lineterminator="\n")
    k = r["kpis"]
    return {"kpis": {chave: k[chave] for chave in (
        "linhas_brutas", "linhas_validas", "inconsistencias", "expurgados", "anomalias", "clientes_unicos",
        "mediana_global_seg", "tempo_ideal_q1_global_seg", "gap_mediana_q1_global_seg", "outliers_boxplot")},
            "motivos": r["inconsistencias"]["Motivo_Inconsistencia"].tolist()}


def main():
    warnings.simplefilter("ignore")
    esperado = {}
    for gerar in (planilha_tipos, planilha_esparsa, planilha_sem_dimensao, planilha_dimensao_errada,
                  planilha_dimensao_maior,
                  planilha_1904, planilha_abas, planilha_sem_estilos, planilha_macro,
                  planilha_base_sintetica, planilha_codigos):
        caminho = gerar()
        esperado[caminho.name] = referencia(caminho)
    for nome in ("base_sintetica", "codigos"):
        esperado[f"{nome}.xlsx"]["processamento"] = saidas_do_processamento(PASTA / f"{nome}.xlsx", nome)
    (PASTA / "esperado.json").write_text(json.dumps(esperado, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"{len(esperado)} planilhas em {PASTA}")


if __name__ == "__main__":
    main()
