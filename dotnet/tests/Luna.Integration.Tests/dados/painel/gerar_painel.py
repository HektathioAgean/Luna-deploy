"""
Referência Python do painel do cliente (PR11).

    .venv\\Scripts\\python dotnet\\tests\\Luna.Integration.Tests\\dados\\painel\\gerar_painel.py

Escreve as bases dirigidas (painel_a.csv, painel_b.xlsx, painel_c.csv), roda
executar() + montar_painel_cliente() (src/pipeline.py) em cada caso e grava em
esperado.json o painel em forma canônica: resumo, janela, tabela e os dados
dos dois gráficos (traços e formas da figura plotly). Grava também casos de
funções isoladas: normalizar_numero_texto, formatar_numero, média do pandas e a
ordem do sort_values em datetime64.

As bases também servem ao fuzz (luna-fuzz-harness, gen_painel.py importa
escrever_base daqui). Nenhum código do motor muda: este arquivo só usa a referência.
"""

from __future__ import annotations

import csv
import datetime as dt
import json
import random
import sys
import warnings
from pathlib import Path

PASTA = Path(__file__).resolve().parent
RAIZ = PASTA.parents[4]
sys.path.insert(0, str(RAIZ))

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402

from src.pipeline import Parametros, executar, montar_painel_cliente  # noqa: E402
from src.utils import formatar_numero, normalizar_numero_texto  # noqa: E402

VOLUMES_TEXTO = ["10,5", "1.234,56", "1,234.56", "7", "0", "0,0", "", "abc", "R$ 12,30", " 3 ", "-2,5", "1.2.3",
                 "1,2,3", "nan", "None", "12.", ".5", "-", "1\xa0234,5", "5e3", "-.75"]


def escrever_base(destino: Path, semente: int, formato: str, volume: str | None, fuso: str = "", n_clientes: int = 6) -> None:
    """Base com empates de horário, janelas na virada do dia e volumes em vários formatos."""
    rng = random.Random(semente)
    clientes = ["00123", "123", "C7"] + [f"K{i}" for i in range(n_clientes - 3)]
    linhas = []
    inicio = dt.datetime(2026, 3, 2, 0, 0)
    for cliente in clientes:
        qtd = rng.choice([1, 2, 3, 5, 15, 16, 17, 40, 90])
        horas = rng.choice([[23, 0], [9, 10], [22, 23, 0, 1], [12]])
        empates = rng.random() < 0.5
        for k in range(qtd):
            dia = rng.randint(0, 40) if not empates else rng.randint(0, 3)
            chegada = inicio + dt.timedelta(days=dia, hours=rng.choice(horas), minutes=rng.choice([0, 15, 30, 59]),
                                            seconds=rng.choice([0, 0, 30]))
            fim = chegada + dt.timedelta(seconds=rng.choice([300, 600, 900, 1800, 2400, 3600, 7200, 20000]))
            if formato == "xlsx" and volume is not None:
                vol = rng.choice([rng.randint(0, 60), None, rng.randint(0, 60)])
            else:
                vol = rng.choice(VOLUMES_TEXTO + [str(rng.randint(1, 99))] * 6)
            linhas.append((cliente, chegada, fim, f"T{rng.randint(1, 9)}", vol))
    rng.shuffle(linhas)
    cab = ["Cod_Cliente", "Chegou_em", "Finalizada_em", "tour_display_id"] + ([volume] if volume else [])
    if formato == "xlsx":
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("Base")
        ws.append(cab)
        for c, a, b, t, v in linhas:
            ws.append([c, a, b, t] + ([v] if volume else []))
        wb.save(destino)
        return
    with open(destino, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, delimiter=";", lineterminator="\n")
        w.writerow(cab)
        for c, a, b, t, v in linhas:
            texto_a = a.strftime("%Y-%m-%dT%H:%M:%S") + fuso
            texto_b = b.strftime("%Y-%m-%dT%H:%M:%S") + fuso
            w.writerow([c, texto_a, texto_b, t] + ([v] if volume else []))


def canon(v):
    if v is None or v is pd.NaT or v is pd.NA:
        return ["nulo"]
    if isinstance(v, (bool, np.bool_)):
        return ["bool", str(bool(v))]
    if isinstance(v, (int, np.integer)):
        return ["int", str(int(v))]
    if isinstance(v, (float, np.floating)):
        return ["nulo"] if np.isnan(v) else ["float", repr(float(v))]
    if isinstance(v, str):
        return ["str", v]
    raise TypeError(type(v))


def lista(valores):
    return [canon(v) for v in (list(valores) if valores is not None else [])]


def grafico(fig) -> dict:
    barras, tempo = fig.data[0], fig.data[1]
    referencias = {t.name.split(" (")[0]: {"nome": t.name, "valor": canon(t.y[0])} for t in fig.data[2:]}
    eixo = fig.layout.xaxis
    return {
        "x": lista(barras.x), "marcas_x": lista(eixo.tickvals), "marcas_texto": lista(eixo.ticktext),
        "faixa": lista(eixo.range),
        "volume": lista(barras.y), "volume_texto": lista(barras.text),
        "volume_detalhe": [lista(par) for par in barras.customdata],
        "tempo": lista(tempo.y), "tempo_texto": None if tempo.text is None else lista(tempo.text),
        "tempo_detalhe": [lista(trio) for trio in tempo.customdata],
        "mediana": referencias.get("Mediana"), "q1": referencias.get("Tempo ideal Q1"),
    }


def aberturas(fig) -> dict:
    pontos = fig.data[0] if fig.data else None
    # Bordas como float: o plotly guarda o 0 e o 1440 da virada do dia como int.
    faixas = [[canon(float(s.x0)), canon(float(s.x1))] for s in fig.layout.shapes if s.type == "rect"]
    linhas = [canon(float(s.x0)) for s in fig.layout.shapes if s.type == "line"]
    return {
        "x": lista(pontos.x) if pontos else [], "y": lista(pontos.y) if pontos else [],
        "texto": lista(pontos.text) if pontos else [],
        "detalhe": [lista(par) for par in pontos.customdata] if pontos else [],
        "faixas": faixas, "linhas": linhas,
    }


def painel_canonico(p: dict) -> dict:
    out = {"cliente": p["cliente"], "qtd_entregas": p["qtd_entregas"], "vazio": p["vazio"]}
    if "dia_semana" in p:
        out["dia_semana"] = p["dia_semana"]
    if "resumo" in p:
        out["resumo"] = {k: canon(v) for k, v in p["resumo"].items()}
    if p.get("vazio"):
        return out
    out["janela"] = None if p["janela"] is None else {k: canon(v) for k, v in p["janela"].items()}
    tabela = p["tabela"]
    out["tabela"] = {"colunas": [str(c) for c in tabela.columns],
                     "linhas": [lista(r) for r in tabela.astype(object).values.tolist()]}
    out["grafico"] = grafico(p["grafico"])
    out["grafico_aberturas"] = None if p["grafico_aberturas"] is None else aberturas(p["grafico_aberturas"])
    return out


def caminho(base: str) -> Path:
    return RAIZ / base if base.startswith("tests/") else PASTA.parent / base


def casos_de_painel() -> list:
    dias = ["TODOS", "SEG", "TER", "QUA", "QUI", "SEX", "SAB", "DOM", " seg ", "todos", "XYZ", ""]
    planos = [
        ("tests/dados/TESTE_data.csv", {}, ["C001", "C002", " C003 ", "C009", "C999", "c001"]),
        ("painel/painel_a.csv", {}, ["00123", "123", "C7", "K1", "K2", "0123"]),
        ("painel/painel_a.csv", {"base_janela": "Finalizada_em", "cobertura_janela": 50, "ajuste_percentual": 10.0,
                                 "eventos_previos": 4, "minimo_apontamentos": 2}, ["00123", "C7", "K2"]),
        ("painel/painel_b.xlsx", {}, ["00123", "123", "C7", "K1"]),
        ("painel/painel_b.xlsx", {"exibir_janelas_entrega": False}, ["C7", "K2"]),
        ("painel/painel_c.csv", {}, ["00123", "C7", "K1", "K2"]),
        ("saidas/esquema_invalido.csv", {}, ["A"]),
    ]
    rng = random.Random(7)
    casos = []
    for base, params, clientes in planos:
        for cliente in clientes:
            for dia in (["TODOS"] + rng.sample(dias, 3)):
                casos.append((base, params, cliente, dia, rng.random() < 0.8))
    return casos


def casos_de_dominio() -> dict:
    rng = random.Random(11)
    textos = VOLUMES_TEXTO + ["1.000.000", "1,000,000", "-", "--5", "1-2", "3,", ",5", "١٢", "12,345.678,9", "  ", "0.000"]
    numeros = [0.0, -0.0, 0.125, 0.135, 2.675, 1234.5, -1234.565, 1e15 + 0.5, 1e-7, 999.995, 123456789.005, 7.0]
    medias = []
    for n in [1, 7, 8, 9, 100, 128, 129, 1000]:
        valores = [rng.uniform(0, 500) * rng.choice([1, 1e-3, 1e5]) for _ in range(n)]
        medias.append({"valores": [repr(v) for v in valores], "media": repr(float(pd.Series(valores).mean()))})
    ordens = []
    for n in [5, 16, 17, 40, 300]:
        chaves = [rng.randrange(rng.choice([2, 5, 1000])) for _ in range(n)]
        serie = pd.DataFrame({"t": pd.to_datetime(np.array(chaves, dtype=np.int64) * 60, unit="s")})
        ordens.append({"chaves": chaves, "ordem": [int(i) for i in serie.sort_values("t").index]})
    return {
        "numero_de_texto": [[t, canon(normalizar_numero_texto(t))] for t in textos],
        "formatar_numero": [[repr(v), formatar_numero(v, 2)] for v in numeros],
        "media": medias,
        "ordem_datetime64": ordens,
    }


def main() -> None:
    warnings.simplefilter("ignore")
    escrever_base(PASTA / "painel_a.csv", 1, "csv", "Vol_caixas")
    escrever_base(PASTA / "painel_b.xlsx", 2, "xlsx", "Volume de caixas")
    escrever_base(PASTA / "painel_c.csv", 3, "csv", None, fuso="-03:00")
    resultados = {}
    paineis = []
    for base, params, cliente, dia, rotulos in casos_de_painel():
        chave = (base, json.dumps(params, sort_keys=True))
        if chave not in resultados:
            resultados[chave] = executar(Parametros(caminho=str(caminho(base)), **params), usar_cache=False)
        p = montar_painel_cliente(resultados[chave], cliente=cliente, dia_semana=dia, mostrar_rotulos_tempo=rotulos)
        paineis.append({"base": base, "params": params, "cliente": cliente, "dia": dia, "rotulos": rotulos,
                        "painel": painel_canonico(p)})
    from src.pipeline import clientes_disponiveis
    disponiveis = {base: clientes_disponiveis(r) for (base, _), r in resultados.items()}
    saida = {"dominio": casos_de_dominio(), "clientes_disponiveis": disponiveis, "paineis": paineis}
    linhas = ",\n".join(json.dumps(c, ensure_ascii=False) for c in paineis)
    cabeca = json.dumps({k: v for k, v in saida.items() if k != "paineis"}, ensure_ascii=False)
    (PASTA / "esperado.json").write_text(cabeca[:-1] + ',\n"paineis": [\n' + linhas + "\n]}\n", encoding="utf-8")
    print(f"{len(paineis)} painéis em {PASTA}")


if __name__ == "__main__":
    main()
